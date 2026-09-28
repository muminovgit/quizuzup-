"""Parses the quiz .xlsx template into a list of question dicts."""

import io

from openpyxl import load_workbook

EXPECTED_HEADERS = [
    "Question",
    "Option A",
    "Option B",
    "Option C",
    "Option D",
    "Correct Answer",
    "Explanation",
]

ANSWER_TO_INDEX = {"A": 0, "B": 1, "C": 2, "D": 3}


class ExcelParseError(Exception):
    pass


def parse_quiz_excel(file_bytes: bytes) -> list[dict]:
    """Parse an uploaded .xlsx file into a list of question dicts.

    Each dict has: question, option_a, option_b, option_c, option_d,
    correct_option (0-3), explanation.
    Raises ExcelParseError with a human-readable message on any problem.
    """
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ExcelParseError(f"Couldn't read the file as an Excel workbook: {exc}") from exc

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ExcelParseError("The sheet is empty.")

    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    for i, expected in enumerate(EXPECTED_HEADERS):
        if i >= len(header) or header[i] != expected:
            raise ExcelParseError(
                "Header row doesn't match the template. Expected columns: "
                + ", ".join(EXPECTED_HEADERS)
            )

    questions = []
    for row_number, row in enumerate(rows[1:], start=2):
        if row is None or all(cell is None for cell in row):
            continue

        question_text = _cell_str(row, 0)
        option_a = _cell_str(row, 1)
        option_b = _cell_str(row, 2)
        option_c = _cell_str(row, 3)
        option_d = _cell_str(row, 4)
        correct_raw = _cell_str(row, 5)
        explanation = _cell_str(row, 6)

        if not question_text:
            raise ExcelParseError(f"Row {row_number}: Question is empty.")
        if not option_a or not option_b:
            raise ExcelParseError(f"Row {row_number}: Option A and Option B are required.")

        correct_letter = correct_raw.strip().upper() if correct_raw else ""
        if correct_letter not in ANSWER_TO_INDEX:
            raise ExcelParseError(
                f"Row {row_number}: Correct Answer must be A, B, C or D (got {correct_raw!r})."
            )
        correct_index = ANSWER_TO_INDEX[correct_letter]

        available_options = [option_a, option_b, option_c, option_d]
        if correct_index >= len(available_options) or not available_options[correct_index]:
            raise ExcelParseError(
                f"Row {row_number}: Correct Answer points to an option that is empty."
            )

        questions.append(
            {
                "question": question_text,
                "option_a": option_a,
                "option_b": option_b,
                "option_c": option_c or None,
                "option_d": option_d or None,
                "correct_option": correct_index,
                "explanation": explanation or None,
            }
        )

    if not questions:
        raise ExcelParseError("No valid questions found in the file.")

    return questions


def _cell_str(row: tuple, index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()
