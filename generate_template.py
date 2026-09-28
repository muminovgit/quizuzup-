"""One-off script that (re)generates quiz_template.xlsx. Not needed to run the bot."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADERS = [
    "Question",
    "Option A",
    "Option B",
    "Option C",
    "Option D",
    "Correct Answer",
    "Explanation",
]

SAMPLE_ROWS = [
    (
        "What is the capital of France?",
        "London",
        "Paris",
        "Berlin",
        "Madrid",
        "B",
        "Paris has been the capital of France since 987 AD.",
    ),
    (
        "Which planet is known as the Red Planet?",
        "Venus",
        "Mars",
        "Jupiter",
        "Saturn",
        "B",
        "Mars appears red due to iron oxide on its surface.",
    ),
    (
        "What is 9 x 9?",
        "72",
        "81",
        "90",
        "99",
        "B",
        "",
    ),
]


def build_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quiz"

    sheet.append(HEADERS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    for row in SAMPLE_ROWS:
        sheet.append(row)

    widths = [40, 18, 18, 18, 18, 14, 40]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = width

    return workbook


if __name__ == "__main__":
    build_workbook().save("quiz_template.xlsx")
    print("quiz_template.xlsx written")
